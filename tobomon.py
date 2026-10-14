"""Tra tên từ Table.xlsx và đánh dấu kết quả trong bản sao của Blank.xls."""

from __future__ import annotations

import argparse
import logging
import os
import re
import tempfile
import unicodedata
from pathlib import Path

import openpyxl
import xlrd
import xlwt


LOGGER = logging.getLogger(__name__)
TONE_MARKS = frozenset("\u0300\u0301\u0303\u0309\u0323")
VOWELS = frozenset("aeiouy")
NAME_SEPARATOR = re.compile(r"\s+-\s+")


def normalize_name(value: object) -> str:
    """Chuẩn hóa cách viết dấu, khoảng trắng và dấu câu của tên."""
    decomposed = unicodedata.normalize("NFD", str(value or "").casefold())
    normalized = []
    vowel_group = []

    def flush_vowel_group() -> None:
        if vowel_group:
            tone_marks = [char for char in vowel_group if char in TONE_MARKS]
            normalized.extend(char for char in vowel_group if char not in TONE_MARKS)
            normalized.extend(tone_marks)
            vowel_group.clear()

    for character in decomposed:
        if character in TONE_MARKS:
            if vowel_group:
                vowel_group.append(character)
        elif character.isspace():
            flush_vowel_group()
        elif character in VOWELS:
            vowel_group.append(character)
        elif unicodedata.category(character).startswith("M") and vowel_group:
            vowel_group.append(character)
        elif character.isalnum():
            flush_vowel_group()
            normalized.append(character)
    flush_vowel_group()
    return "".join(normalized)


def teacher_name(value: object) -> str:
    """Lấy phần tên, bỏ tài khoản đứng sau dấu gạch ngang trong mẫu."""
    return NAME_SEPARATOR.split(str(value or "").strip(), maxsplit=1)[0]


def copy_style(source_book, style_index: int):
    """Chuyển định dạng ô từ xlrd sang xlwt."""
    source_xf = source_book.xf_list[style_index]
    style = xlwt.XFStyle()

    source_font = source_book.font_list[source_xf.font_index]
    font = xlwt.Font()
    for attribute in (
        "name", "height", "bold", "italic", "underline", "struck_out",
        "colour_index", "escapement", "family", "charset", "outline",
        "shadow", "character_spacing", "word_spacing",
    ):
        if hasattr(source_font, attribute):
            setattr(font, attribute, getattr(source_font, attribute))
    style.font = font

    source_alignment = source_xf.alignment
    alignment = xlwt.Alignment()
    for attribute in ("horz", "vert", "wrap", "rotation", "shrink_to_fit", "indent"):
        if hasattr(source_alignment, attribute):
            setattr(alignment, attribute, getattr(source_alignment, attribute))
    style.alignment = alignment

    source_border = source_xf.border
    borders = xlwt.Borders()
    for side in ("left", "right", "top", "bottom", "diag"):
        source_side = getattr(source_border, side, None)
        if source_side is not None:
            destination_side = getattr(borders, side)
            destination_side.line_style = getattr(source_side, "line_style", 0)
            destination_side.colour_index = getattr(source_side, "colour_index", 0)
    style.borders = borders

    source_background = source_xf.background
    pattern = xlwt.Pattern()
    pattern.pattern = source_background.fill_pattern
    pattern.pattern_fore_colour = source_background.pattern_colour_index
    pattern.pattern_back_colour = source_background.background_colour_index
    style.pattern = pattern

    number_format = source_book.format_map.get(source_xf.format_key)
    style.num_format_str = number_format.format_str if number_format else "General"
    return style


def copy_template(template_path: Path):
    """Tạo workbook XLS mới từ nội dung, định dạng và kích thước trong mẫu."""
    source_book = xlrd.open_workbook(str(template_path), formatting_info=True)
    output_book = xlwt.Workbook()
    style_cache = {}

    for source_sheet in source_book.sheets():
        output_sheet = output_book.add_sheet(
            source_sheet.name, cell_overwrite_ok=True
        )
        for row_index in range(source_sheet.nrows):
            row_info = source_sheet.rowinfo_map.get(row_index)
            if row_info:
                output_sheet.row(row_index).height = row_info.height
                output_sheet.row(row_index).hidden = row_info.hidden

            for column_index in range(source_sheet.ncols):
                cell = source_sheet.cell(row_index, column_index)
                if cell.xf_index not in style_cache:
                    style_cache[cell.xf_index] = copy_style(
                        source_book, cell.xf_index
                    )
                output_sheet.write(
                    row_index,
                    column_index,
                    cell.value,
                    style_cache[cell.xf_index],
                )

        for column_index, column_info in source_sheet.colinfo_map.items():
            output_sheet.col(column_index).width = column_info.width
            output_sheet.col(column_index).hidden = column_info.hidden

    return source_book, output_book


def map_names(input_path: Path, template_path: Path, output_path: Path) -> int:
    """Đánh dấu cột C của mẫu cho các tên tìm thấy trong file đầu vào."""
    input_book = openpyxl.load_workbook(
        input_path, read_only=True, data_only=True
    )
    try:
        if not input_book.worksheets:
            raise ValueError(f"File đầu vào không có trang tính: {input_path}")
        input_sheet = input_book.worksheets[0]
        input_names = [
            input_sheet.cell(row_index, 2).value
            for row_index in range(2, input_sheet.max_row + 1)
        ]
    finally:
        input_book.close()

    source_book, output_book = copy_template(template_path)
    if not source_book.sheets():
        raise ValueError(f"File mẫu không có trang tính: {template_path}")
    source_sheet = source_book.sheet_by_index(0)
    output_sheet = output_book.get_sheet(source_sheet.name)

    names_by_key = {}
    for row_index in range(1, source_sheet.nrows):
        name = teacher_name(source_sheet.cell_value(row_index, 1))
        key = normalize_name(name)
        if key:
            names_by_key.setdefault(key, []).append(row_index)

    matched_rows = set()
    unmatched_names = []
    for value in input_names:
        if value is None or not str(value).strip():
            continue

        key = normalize_name(value)
        matching_rows = names_by_key.get(key, [])
        if matching_rows:
            matched_rows.update(matching_rows)
        else:
            unmatched_names.append(str(value).strip())
            LOGGER.warning("Không tìm thấy tên tương tự: %s", str(value).strip())

    for row_index in matched_rows:
        output_sheet.write(
            row_index,
            2,
            1,
            copy_style(source_book, source_sheet.cell(row_index, 2).xf_index),
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(
            delete=False, suffix=".xls", dir=output_path.parent
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        output_book.save(str(temporary_path))
        os.replace(temporary_path, output_path)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()

    LOGGER.info(
        "Đã đánh dấu %d dòng; có %d tên không tìm thấy. File kết quả: %s",
        len(matched_rows),
        len(unmatched_names),
        output_path,
    )
    return len(matched_rows)


def main() -> None:
    base_directory = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=base_directory / "Table.xlsx")
    parser.add_argument("--template", type=Path, default=base_directory / "Blank.xls")
    parser.add_argument("--output", type=Path, default=base_directory / "To_x.xls")
    arguments = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    map_names(arguments.input, arguments.template, arguments.output)


if __name__ == "__main__":
    main()